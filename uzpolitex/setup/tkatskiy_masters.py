"""Ткацкий etapi ma'lumotlari: bo'lim, 143 xodim (Сотрудники listidan), normalar.

Bir martalik: bench --site uz.politex execute uzpolitex.setup.tkatskiy_masters.setup
Idempotent — mavjudlarini o'tkazib yuboradi.

DIQQAT: xodimlarda tug'ilgan sana/ishga kirgan sana Excelda YO'Q — placeholder
qo'yiladi (1990-01-01 / 2025-06-01), HR keyin to'ldiradi (ekstruderdagidek).
"""

import openpyxl

import frappe

XLSX = (
	"/tmp/claude-1000/-home-user-frappe-bench-16/"
	"b4e439fa-33f2-4f38-875d-903c13506dc6/scratchpad/uzpolitex_sheet.xlsx"
)
COMPANY = "Uzpolitex Holding"
DEPT = "Ткацкий - UH"

# Норма listi (D:F ustunlar, g/m): Лента-45/Лента-55 uchun Excelda norma yo'q
NORMALAR = {
	"Лента-50": (25, 31),
	"Рулон-35": (37, 41),
	"Рулон-40": (42, 46),
	"Рулон-45": (49, 53),
	"Рулон-50": (54, 58),  # Excelda 'Рулон-50 ' (probel bilan) — tozalangan
	"Рулон-55": (59, 62),
}


def setup():
	_department()
	_normalar()
	_xodimlar()
	frappe.db.commit()


def _department():
	if not frappe.db.exists("Department", DEPT):
		frappe.get_doc(
			{
				"doctype": "Department",
				"department_name": "Ткацкий",
				"company": COMPANY,
				"parent_department": "Production - UH",
			}
		).insert()
		print(f"Department yaratildi: {DEPT}")


def _normalar():
	for item, (mn, mx) in NORMALAR.items():
		if frappe.db.exists("Item", item):
			frappe.db.set_value("Item", item, {"uz_norma_min": mn, "uz_norma_max": mx})
			print(f"Norma: {item} = {mn}–{mx} г/м")
		else:
			print(f"DIQQAT: Item topilmadi: {item}")


def _xodimlar():
	wb = openpyxl.load_workbook(XLSX, data_only=True)
	ws = wb["Сотрудники"]
	mavjud = set(frappe.get_all("Employee", pluck="employee_name"))
	yaratildi = 0
	for r in range(3, ws.max_row + 1):
		fio = ws.cell(row=r, column=3).value
		bolim = ws.cell(row=r, column=6).value
		if not (isinstance(fio, str) and fio.strip() and bolim == "Ткацкий"):
			continue
		fio = fio.strip()
		if fio in mavjud:
			continue
		frappe.get_doc(
			{
				"doctype": "Employee",
				"first_name": fio,
				"employee_name": fio,
				"gender": "Prefer not to say",  # PLACEHOLDER — HR to'ldiradi
				"date_of_birth": "1990-01-01",  # PLACEHOLDER
				"date_of_joining": "2025-06-01",  # PLACEHOLDER
				"company": COMPANY,
				"department": DEPT,
				"status": "Active",
			}
		).insert()
		yaratildi += 1
	print(f"Ткацкий xodimlari: {yaratildi} ta yaratildi")
