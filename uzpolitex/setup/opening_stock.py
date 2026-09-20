"""Google Sheets «Склад» listidan boshlang'ich qoldiqlarni yuklash (2026-09-06).

Bir martalik: bench --site uz.politex execute uzpolitex.setup.opening_stock.run

Manba bloklari:
- B:E  Сырьевой склад            → "Сырьевой склад - UH"
- L:P  Sexlar bo'yicha (Отдел)   → "<Отдел> - UH", Готовый продукт → "ГП склад - UH"
  (G:I "Общий склад" bloki L:P ning jamlanmasi — takror yuklanmaydi)

Narxlar: Приход listidan har itemning OXIRGI $ narxi (Сумма $ / Кол-во).
Narxsiz itemlar (полуфабрикат/ГП) allow_zero_valuation_rate bilan kiradi —
tannarx keyin aniqlanganda Stock Reconciliation bilan to'g'rilanadi.

Chetda qoladiganlar (ERPNext manfiy qoldiqni qabul qilmaydi, #DIV/0! noma'lum):
ro'yxati oxirida chop etiladi — Excelda tuzatilishi kerak.
"""

import openpyxl

import frappe
from frappe.utils import flt, nowdate

XLSX = (
	"/tmp/claude-1000/-home-user-frappe-bench-16/"
	"b4e439fa-33f2-4f38-875d-903c13506dc6/scratchpad/uzpolitex_sheet.xlsx"
)
COMPANY = "Uzpolitex Holding"
SIRYO_WH = "Сырьевой склад - UH"
GP_WH = "ГП склад - UH"
DEPT_WH = {
	"Экструдор": "Экструдор - UH",
	"Ткацкий": "Ткацкий - UH",
	"Ламинант": "Ламинант - UH",
	"Печать": "Печать - UH",
	"Конверт": "Конверт - UH",
}
UOM_MAP = {"КГ": "Kg", "ШТ": "Unit", "МЕТР": "Meter"}


def _cell(ws, r, c):
	v = ws.cell(row=r, column=c).value
	return v.strip() if isinstance(v, str) else v


def _last_prices(wb):
	"""Приход: item → oxirgi $ narx (Сумма $ / Кол-во).

	Nom D yoki E ustunida bo'lishi mumkin, ikkinchisida izoh matn turadi
	("41 та келди") — shuning uchun faqat DB'dagi Item nomiga mos kelgani olinadi.
	"""
	items = set(frappe.get_all("Item", pluck="name"))
	ws = wb["Приход"]
	best = {}  # name -> (date, rate)
	for r in range(3, ws.max_row + 1):
		d, e = _cell(ws, r, 4), _cell(ws, r, 5)
		name = d if d in items else (e if e in items else None)
		qty, usd, date = _cell(ws, r, 6), _cell(ws, r, 13), _cell(ws, r, 2)
		if (
			isinstance(name, str)
			and name
			and isinstance(qty, (int, float))
			and isinstance(usd, (int, float))
			and qty > 0
			and usd > 0
		):
			key = str(date or "")
			if name not in best or key >= best[name][0]:
				best[name] = (key, usd / qty)
	return {k: v[1] for k, v in best.items()}


def _collect(wb):
	"""Склад listidan (item, warehouse, qty, uom, tip) qatorlari + chetlatilganlar."""
	ws = wb["Склад"]
	entries, skipped = [], []

	def add(name, uom, qty, tip, wh):
		if not isinstance(qty, (int, float)):
			skipped.append((name, wh, qty, "raqam emas (#DIV/0!?)"))
		elif qty < 0:
			skipped.append((name, wh, qty, "manfiy qoldiq"))
		elif qty > 0:
			entries.append((name, wh, flt(qty, 3), uom, tip))

	for r in range(3, ws.max_row + 1):
		# Сырьевой склад bloki
		name = _cell(ws, r, 2)
		if isinstance(name, str) and name:
			add(name, _cell(ws, r, 3), _cell(ws, r, 4), _cell(ws, r, 5), SIRYO_WH)
		# Otdel bloki
		name = _cell(ws, r, 12)
		if isinstance(name, str) and name:
			tip, otdel = _cell(ws, r, 15), _cell(ws, r, 16)
			if tip == "Готовый продукт":
				wh = GP_WH
			elif otdel in DEPT_WH:
				wh = DEPT_WH[otdel]
			else:
				skipped.append((name, otdel, _cell(ws, r, 14), "otdel noma'lum"))
				continue
			add(name, _cell(ws, r, 13), _cell(ws, r, 14), tip, wh)

	return entries, skipped


def _ensure_item(name, uom, tip, otdel_wh, skipped):
	if frappe.db.exists("Item", name):
		return True
	group = None
	for g in (f"{tip} - {otdel_wh}", tip):
		if frappe.db.exists("Item Group", g):
			group = g
			break
	if not group:
		skipped.append((name, otdel_wh, "-", "item ham, mos guruh ham yo'q"))
		return False
	frappe.get_doc(
		{
			"doctype": "Item",
			"item_code": name,
			"item_group": group,
			"stock_uom": UOM_MAP.get(uom, "Kg"),
			"is_stock_item": 1,
		}
	).insert()
	return True


def run():
	wb = openpyxl.load_workbook(XLSX, data_only=True)
	prices = _last_prices(wb)
	entries, skipped = _collect(wb)

	expense_account = frappe.db.get_value(
		"Account", {"account_type": "Temporary", "company": COMPANY, "is_group": 0}
	)
	if not expense_account:
		frappe.throw("Temporary Opening account topilmadi")

	sr = frappe.new_doc("Stock Reconciliation")
	sr.purpose = "Opening Stock"
	sr.company = COMPANY
	sr.posting_date = nowdate()
	sr.set_posting_time = 1
	sr.expense_account = expense_account

	no_price = []
	for name, wh, qty, uom, tip in entries:
		wh_dept = wh.replace(" - UH", "")
		if not _ensure_item(name, uom, tip, wh_dept, skipped):
			continue
		row = {"item_code": name, "warehouse": wh, "qty": qty}
		rate = prices.get(name)
		# faqat siryo/rasxodnikka xarid narxi ishonchli; qolgani 0 (keyin to'g'rilanadi)
		if rate and tip in ("Сырьё", "Расходник"):
			row["valuation_rate"] = flt(rate, 4)
		else:
			row["allow_zero_valuation_rate"] = 1
			no_price.append(f"{name} ({wh})")
		sr.append("items", row)

	sr.insert()
	sr.submit()
	frappe.db.commit()

	print(f"OK: {sr.name} — {len(sr.items)} qator yuklandi ({sr.posting_date})")
	print(f"Narxsiz (qiymat=0) qatorlar: {len(no_price)}")
	for x in no_price:
		print("   0$:", x)
	print(f"Chetda qolganlar: {len(skipped)}")
	for name, wh, qty, why in skipped:
		print(f"   SKIP: {name} | {wh} | {qty} | {why}")
