"""Ламинат etapi ma'lumotlari: bo'lim, 8 xodim, В 1М(ГР) normalari.

Bir martalik: bench --site uz.politex execute uzpolitex.setup.laminat_masters.setup
Idempotent. Xodim sanalari PLACEHOLDER (HR to'ldiradi).
"""

import openpyxl

import frappe

XLSX = (
	"/tmp/claude-1000/-home-user-frappe-bench-16/"
	"b4e439fa-33f2-4f38-875d-903c13506dc6/scratchpad/uzpolitex_sheet.xlsx"
)
COMPANY = "Uzpolitex Holding"
DEPT = "Ламинант - UH"

# Норма listi H:J (g/m); Ламинат Лента-55 va Лента рулон 45 uchun Excelda norma YO'Q
NORMALAR = {
	"Ламинат Рулон-35": (52, 60),
	"Ламинат Рулон-40": (60, 65),
	"Ламинат Рулон-45": (68, 76),
	"Ламинат Рулон-50": (75, 80),
	"Ламинат Рулон-55": (80, 84),
	"Ламинат Лента рулон 50": (50, 55),
}


def setup():
	if not frappe.db.exists("Department", DEPT):
		frappe.get_doc(
			{
				"doctype": "Department",
				"department_name": "Ламинант",
				"company": COMPANY,
				"parent_department": "Production - UH",
			}
		).insert()
		print(f"Department yaratildi: {DEPT}")

	for item, (mn, mx) in NORMALAR.items():
		if frappe.db.exists("Item", item):
			frappe.db.set_value("Item", item, {"uz_norma_min": mn, "uz_norma_max": mx})
			print(f"Norma: {item} = {mn}–{mx} г/м")
		else:
			print(f"DIQQAT: Item topilmadi: {item}")

	wb = openpyxl.load_workbook(XLSX, data_only=True)
	ws = wb["Сотрудники"]
	mavjud = set(frappe.get_all("Employee", pluck="employee_name"))
	yaratildi = 0
	for r in range(3, ws.max_row + 1):
		fio = ws.cell(row=r, column=3).value
		bolim = ws.cell(row=r, column=6).value
		if not (isinstance(fio, str) and fio.strip() and bolim == "Ламинант"):
			continue
		fio = fio.strip()
		if fio in mavjud:
			continue
		frappe.get_doc(
			{
				"doctype": "Employee",
				"first_name": fio,
				"employee_name": fio,
				"gender": "Prefer not to say",  # PLACEHOLDER — HR to'ldiradi
				"date_of_birth": "1990-01-01",  # PLACEHOLDER
				"date_of_joining": "2025-06-01",  # PLACEHOLDER
				"company": COMPANY,
				"department": DEPT,
				"status": "Active",
			}
		).insert()
		yaratildi += 1
	print(f"Ламинант xodimlari: {yaratildi} ta yaratildi")
	frappe.db.commit()
